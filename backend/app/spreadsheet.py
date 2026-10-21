"""Leitura de planilhas com suporte a varredura em fluxo para arquivos XLSX."""

import io
import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


def nome_coluna_excel(indice: int) -> str:
    nome = ""
    while indice >= 0:
        indice, resto = divmod(indice, 26)
        nome = chr(65 + resto) + nome
        indice -= 1
    return f"Coluna {nome}"


def normalizar_cabecalhos(valores) -> list[str]:
    nomes = []
    usados = set()
    for indice, valor in enumerate(valores):
        nome = str(valor).strip() if valor is not None and not pd.isna(valor) else ""
        if not nome:
            nome = nome_coluna_excel(indice)
        base = nome
        sufixo = 2
        while nome.casefold() in usados:
            nome = f"{base} ({sufixo})"
            sufixo += 1
        usados.add(nome.casefold())
        nomes.append(nome)
    return nomes


def detectar_cabecalho(valores) -> bool:
    presentes = [v for v in valores if v is not None and str(v).strip()]
    if not presentes:
        return False
    padrao = re.compile(r"\b(contrato|contract|n[uú]mero|c[oó]digo|id)\b", re.IGNORECASE)
    if any(padrao.search(str(v)) for v in presentes):
        return True
    cara_de_dado = sum(
        bool(re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._/ -]{2,}", str(v).strip()))
        and any(c.isdigit() for c in str(v))
        for v in presentes
    )
    return cara_de_dado < max(1, (len(presentes) + 1) // 2)


def valor_contrato(cell) -> str:
    valor = cell.value
    if valor is None:
        return ""
    if isinstance(valor, str):
        return valor.strip()
    if isinstance(valor, bool):
        return str(valor)
    if isinstance(valor, (int, float)):
        formato = str(getattr(cell, "number_format", ""))
        if isinstance(valor, float) and valor.is_integer():
            valor = int(valor)
        # Conserva zeros à esquerda aplicados como formato numérico simples no Excel.
        if isinstance(valor, int) and re.fullmatch(r"0+", formato):
            return str(valor).zfill(len(formato))
        return str(valor)
    return str(valor).strip()


def _abrir_xlsx(source):
    if isinstance(source, (str, Path)):
        return load_workbook(source, read_only=True, data_only=True)
    return load_workbook(io.BytesIO(source), read_only=True, data_only=True)


def inspecionar_planilha_xlsx(source, tem_cabecalho=None):
    workbook = _abrir_xlsx(source)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows()
        first = next(rows, None)
        if first is None:
            raise ValueError("A planilha enviada está vazia.")
        first_values = [cell.value for cell in first]
        if tem_cabecalho is None:
            tem_cabecalho = detectar_cabecalho(first_values)
        names = normalizar_cabecalhos(first_values) if tem_cabecalho else [nome_coluna_excel(i) for i in range(len(first_values))]
        counts = [0] * len(names)
        rows_with_data = 0
        start_row = 2 if tem_cabecalho else 1
        for row in sheet.iter_rows(min_row=start_row, max_col=len(names)):
            if any(valor_contrato(cell) for cell in row):
                rows_with_data += 1
            for index, cell in enumerate(row):
                if valor_contrato(cell):
                    counts[index] += 1
        return {"colunas": names, "contagens": counts, "linhas_dados": rows_with_data, "tem_cabecalho": tem_cabecalho}
    finally:
        workbook.close()


def extrair_contratos_xlsx(source, coluna_contrato=None, tem_cabecalho=None):
    workbook = _abrir_xlsx(source)
    try:
        sheet = workbook.active
        first = next(sheet.iter_rows(max_row=1), None)
        if first is None:
            raise ValueError("A planilha enviada está vazia.")
        first_values = [cell.value for cell in first]
        if tem_cabecalho is None:
            tem_cabecalho = detectar_cabecalho(first_values)
        names = normalizar_cabecalhos(first_values) if tem_cabecalho else [nome_coluna_excel(i) for i in range(len(first_values))]
        start_row = 2 if tem_cabecalho else 1
        counts = [0] * len(names)
        row_count = 0
        for row in sheet.iter_rows(min_row=start_row, max_col=len(names)):
            if any(valor_contrato(cell) for cell in row):
                row_count += 1
            for index, cell in enumerate(row):
                if valor_contrato(cell):
                    counts[index] += 1

        if coluna_contrato is not None:
            if coluna_contrato not in names:
                raise ValueError(f"A coluna '{coluna_contrato}' não existe na planilha.")
            column_index = names.index(coluna_contrato)
        else:
            candidates = [i for i, count in enumerate(counts) if count]
            if not candidates:
                raise ValueError("Não encontrei contratos preenchidos na planilha.")
            column_index = next((i for i in candidates if "contrato" in names[i].casefold()), candidates[0])
            coluna_contrato = names[column_index]

        contratos = []
        for row in sheet.iter_rows(min_row=start_row, min_col=column_index + 1, max_col=column_index + 1):
            contrato = valor_contrato(row[0])
            if contrato:
                contratos.append(contrato)
        return contratos, row_count, coluna_contrato, tem_cabecalho
    finally:
        workbook.close()


def iterar_contratos_xlsx(source, coluna_contrato, tem_cabecalho):
    workbook = _abrir_xlsx(source)
    try:
        sheet = workbook.active
        first = next(sheet.iter_rows(max_row=1), None)
        if first is None:
            return
        first_values = [cell.value for cell in first]
        names = normalizar_cabecalhos(first_values) if tem_cabecalho else [nome_coluna_excel(i) for i in range(len(first_values))]
        if coluna_contrato not in names:
            raise ValueError(f"A coluna '{coluna_contrato}' não existe na planilha.")
        index = names.index(coluna_contrato) + 1
        start_row = 2 if tem_cabecalho else 1
        for row in sheet.iter_rows(min_row=start_row, min_col=index, max_col=index):
            value = valor_contrato(row[0])
            if value:
                yield value
    finally:
        workbook.close()
