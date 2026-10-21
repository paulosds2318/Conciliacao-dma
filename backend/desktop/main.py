import io
import queue
import sys
import threading
from datetime import datetime
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

import pandas as pd
from app.config import CORAL_API_URL, MODO_TESTE_API
from app.services import detectar_cabecalho, processar_e_enviar_planilha
from app.spreadsheet import inspecionar_planilha_xlsx
from app.journal import connect as connect_journal, count_items, iter_items, page_items

RED = "#f40000"
RED_HOVER = "#d90000"
RED_PALE = "#fff3f3"
INK = "#202020"
MUTED = "#777777"
WHITE = "#ffffff"
BORDER = "#e9e9e9"
GREEN = "#18794e"


def safe_excel_text(value):
    text = "" if value is None else str(value)
    if text.startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + text
    return text


def resource_path(relative_path):
    base = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parents[2]))
    return base / relative_path


class DMADesktopApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Concilia\u00e7\u00e3o DMA | Foco")
        self.geometry("1260x840")
        self.minsize(1080, 720)
        self.configure(bg=WHITE)
        icon_path = resource_path("frontend/foco.ico")
        if icon_path.exists():
            self.iconbitmap(default=str(icon_path))
        self.result_queue = queue.Queue()
        self.inspection_queue = queue.Queue()
        self.selected_file = tk.StringVar()
        self.selected_name = tk.StringVar(value="Nenhuma planilha selecionada")
        self.selected_folder = tk.StringVar(value="Excel .xlsx ou .xls")
        self.action = tk.StringVar(value="Desconsiderar DMA")
        self.contract_column = tk.StringVar()
        self.contract_count = tk.StringVar(value="Contratos encontrados: --")
        self.tem_cabecalho = tk.BooleanVar(value=True)
        self.preview_df = None
        self.column_counts = {}
        self.log_page = 0
        self.page_size = 200
        self.columns = []
        self.log_filter = tk.StringVar(value="Todos")
        self.log_rows = []
        self.details_db_path = None
        self.job_id = None
        self.report_metadata = None
        self._build_styles()
        self._build_layout()

    def _build_styles(self):
        style = ttk.Style(self)
        style.theme_use("clam")
        style.configure("Treeview", rowheight=38, font=("Arial", 10), background=WHITE, fieldbackground=WHITE, foreground=INK, borderwidth=0)
        style.configure("Treeview.Heading", font=("Arial", 10, "bold"), background="#f7f7f7", foreground=INK, relief="flat", padding=(10, 10))
        style.map("Treeview", background=[("selected", RED_PALE)], foreground=[("selected", INK)])
        style.configure("Column.TCombobox", padding=5, font=("Arial", 9))

    def _build_layout(self):
        tk.Frame(self, bg=RED, width=14).place(x=0, y=0, relheight=1)
        body = tk.Frame(self, bg=WHITE, width=1160, height=780)
        body.place(relx=0.5, rely=0.5, anchor="center")
        body.grid_propagate(False)
        body.grid_columnconfigure(0, minsize=320, weight=3)
        body.grid_columnconfigure(1, minsize=740, weight=7)
        body.grid_rowconfigure(0, weight=1)

        intro = tk.Frame(body, bg=WHITE)
        intro.grid(row=0, column=0, sticky="nsew", padx=(0, 58))
        self.logo = None
        logo_path = resource_path("frontend/foco-logo.png")
        if logo_path.exists():
            try:
                self.logo = tk.PhotoImage(file=str(logo_path)).subsample(5, 5)
                tk.Label(intro, image=self.logo, bg=WHITE).place(x=0, y=8, anchor="nw")
            except tk.TclError:
                pass

        copy = tk.Frame(intro, bg=WHITE)
        copy.place(relx=0, rely=0.53, anchor="w", relwidth=1)
        tk.Label(copy, text="OPERA\u00c7\u00c3O  /  DMA CORAL", fg=RED, bg=WHITE, font=("Arial", 10, "bold")).pack(anchor="w")
        tk.Frame(copy, bg=RED, width=42, height=3).pack(anchor="w", pady=(20, 17))
        tk.Label(copy, text="Importa\u00e7\u00e3o de\ncontratos", fg=INK, bg=WHITE, font=("Arial", 30, "bold"), justify="left", anchor="w").pack(anchor="w")
        tk.Label(copy, text="Envie sua planilha e escolha como a regra DMA ser\u00e1 aplicada aos contratos.", fg=MUTED, bg=WHITE, font=("Arial", 11), justify="left", wraplength=280, anchor="w").pack(anchor="w", pady=(12, 0))
        tk.Label(intro, text="Foco Aluguel de Carros\nConcilia\u00e7\u00e3o DMA", fg="#999999", bg=WHITE, font=("Arial", 8), justify="left").place(x=0, rely=1.0, y=-4, anchor="sw")

        right = tk.Frame(body, bg=WHITE)
        right.grid(row=0, column=1, sticky="nsew")
        right.grid_columnconfigure(0, weight=1)
        right.grid_rowconfigure(3, weight=1)

        form = tk.Frame(right, bg=WHITE, highlightbackground=BORDER, highlightthickness=1, padx=22, pady=20)
        form.grid(row=0, column=0, sticky="ew", pady=(8, 0))
        tk.Label(form, text="Preparar processamento", fg=INK, bg=WHITE, font=("Arial", 16, "bold")).pack(anchor="w")
        tk.Label(form, text="Preencha as informa\u00e7\u00f5es abaixo para iniciar o envio.", fg=MUTED, bg=WHITE, font=("Arial", 9)).pack(anchor="w", pady=(5, 15))

        tk.Label(form, text="PLANILHA DE CONTRATOS", fg=INK, bg=WHITE, font=("Arial", 9, "bold")).pack(anchor="w", pady=(0, 7))
        filebox = tk.Frame(form, bg=WHITE, highlightbackground="#dedede", highlightthickness=1, padx=10, pady=8)
        filebox.pack(fill="x")
        file_info = tk.Frame(filebox, bg=WHITE)
        file_info.pack(side="left", fill="x", expand=True, padx=(0, 7))
        tk.Label(file_info, textvariable=self.selected_name, fg=INK, bg=WHITE, font=("Arial", 9, "bold"), anchor="w").pack(fill="x")
        tk.Label(file_info, textvariable=self.selected_folder, fg=MUTED, bg=WHITE, font=("Arial", 8), anchor="w").pack(fill="x", pady=(3, 0))
        actions = tk.Frame(filebox, bg=WHITE)
        actions.pack(side="right")
        self.select_button = tk.Button(actions, text="Selecionar", command=self._choose_file, bg=RED_PALE, fg=RED, activebackground="#ffe5e5", relief="flat", padx=11, pady=6, font=("Arial", 8, "bold"), cursor="hand2")
        self.select_button.pack(side="left")
        self.remove_button = tk.Button(actions, text="Remover", command=self._remove_file, state="disabled", bg=WHITE, fg=MUTED, activebackground="#f5f5f5", relief="flat", padx=8, pady=6, font=("Arial", 8), cursor="hand2")
        self.remove_button.pack(side="left", padx=(3, 0))

        tk.Label(form, text="Coluna do n\u00famero do contrato", fg=MUTED, bg=WHITE, font=("Arial", 8)).pack(anchor="w", pady=(9, 3))
        self.column_box = ttk.Combobox(form, textvariable=self.contract_column, state="disabled", style="Column.TCombobox")
        self.column_box.pack(fill="x")
        self.column_box.bind("<<ComboboxSelected>>", self._update_contract_count)
        self.header_checkbox = tk.Checkbutton(form, text="A primeira linha contém os nomes das colunas", variable=self.tem_cabecalho, command=self._reload_columns, state="disabled", bg=WHITE, fg=MUTED, activebackground=WHITE, selectcolor=WHITE, font=("Arial", 8))
        self.header_checkbox.pack(anchor="w", pady=(3, 0))
        tk.Label(form, textvariable=self.contract_count, fg=RED, bg=WHITE, font=("Arial", 9, "bold")).pack(anchor="w", pady=(4, 0))

        tk.Label(form, text="A\u00c7\u00c3O PARA OS CONTRATOS", fg=INK, bg=WHITE, font=("Arial", 9, "bold")).pack(anchor="w", pady=(14, 7))
        choices = tk.Frame(form, bg=WHITE)
        choices.pack(fill="x")
        self.disregard_radio = tk.Radiobutton(choices, text="Desconsiderar DMA", variable=self.action, value="Desconsiderar DMA", indicatoron=False, anchor="w", padx=10, pady=9, font=("Arial", 9), fg=INK, bg=WHITE, activebackground=RED_PALE, selectcolor=RED_PALE, relief="flat", bd=0, highlightthickness=1, highlightbackground="#dedede", cursor="hand2")
        self.disregard_radio.pack(side="left", fill="x", expand=True, padx=(0, 7))
        self.consider_radio = tk.Radiobutton(choices, text="Considerar DMA", variable=self.action, value="Considerar DMA", indicatoron=False, anchor="w", padx=10, pady=9, font=("Arial", 9), fg=INK, bg=WHITE, activebackground=RED_PALE, selectcolor=RED_PALE, relief="flat", bd=0, highlightthickness=1, highlightbackground="#dedede", cursor="hand2")
        self.consider_radio.pack(side="left", fill="x", expand=True, padx=(7, 0))

        self.run_button = tk.Button(form, text="Processar contratos  \u2192", command=self._process, bg=RED, fg=WHITE, activebackground=RED_HOVER, activeforeground=WHITE, relief="flat", padx=12, pady=10, font=("Arial", 9, "bold"), cursor="hand2", disabledforeground="#ffffff")
        self.run_button.pack(fill="x", pady=(16, 0))
        self.status = tk.StringVar(value="")
        tk.Label(right, textvariable=self.status, fg=MUTED, bg=WHITE, font=("Arial", 9), anchor="w").grid(row=1, column=0, sticky="ew", pady=(7, 4))

        results = tk.Frame(right, bg="#fafafa", highlightbackground=BORDER, highlightthickness=1, padx=13, pady=9)
        results.grid(row=2, column=0, sticky="ew")
        results.grid_columnconfigure(0, weight=1)
        tk.Label(results, text="Resultado do processamento", fg=INK, bg="#fafafa", font=("Arial", 10, "bold")).grid(row=0, column=0, sticky="w")
        metrics = tk.Frame(results, bg="#fafafa")
        metrics.grid(row=1, column=0, sticky="ew", pady=(9, 0))
        self.total_label = self._metric(metrics, "TOTAL", "\u2014", INK)
        self.success_label = self._metric(metrics, "SUCESSOS", "\u2014", GREEN)
        self.fail_label = self._metric(metrics, "FALHAS", "\u2014", RED)
        self.total_label.pack(side="left", padx=(0, 18))
        self.success_label.pack(side="left", padx=(0, 18))
        self.fail_label.pack(side="left")
        self.column_used = tk.StringVar(value="")
        tk.Label(results, textvariable=self.column_used, fg=MUTED, bg="#fafafa", font=("Arial", 8), anchor="w").grid(row=2, column=0, sticky="ew", pady=(7, 0))

        table_panel = tk.Frame(right, bg=WHITE)
        table_panel.grid(row=3, column=0, sticky="nsew", pady=(7, 0))
        table_panel.grid_columnconfigure(0, weight=1)
        table_panel.grid_rowconfigure(1, weight=1)
        filter_bar = tk.Frame(table_panel, bg=WHITE)
        filter_bar.grid(row=0, column=0, columnspan=2, sticky="ew", pady=(0, 5))
        tk.Label(filter_bar, text="Filtrar resultados:", fg=MUTED, bg=WHITE, font=("Arial", 8)).pack(side="left")
        self.log_filter_box = ttk.Combobox(
            filter_bar,
            textvariable=self.log_filter,
            values=("Todos", "Somente falhas", "Somente sucessos"),
            state="readonly",
            width=21,
            style="Column.TCombobox",
        )
        self.log_filter_box.pack(side="left", padx=(7, 0))
        self.log_filter_box.bind("<<ComboboxSelected>>", self._change_log_filter)
        self.page_label = tk.StringVar(value="Página 1")
        self.previous_page_button = tk.Button(filter_bar, text="‹", command=self._previous_page, state="disabled", relief="flat", bg=WHITE, fg=INK)
        self.previous_page_button.pack(side="right", padx=(4, 0))
        tk.Label(filter_bar, textvariable=self.page_label, fg=MUTED, bg=WHITE, font=("Arial", 8)).pack(side="right", padx=4)
        self.next_page_button = tk.Button(filter_bar, text="›", command=self._next_page, state="disabled", relief="flat", bg=WHITE, fg=INK)
        self.next_page_button.pack(side="right", padx=(4, 0))
        self.export_button = tk.Button(
            filter_bar,
            text="Exportar relatório Excel",
            command=self._export_report,
            state="disabled",
            bg=RED_PALE,
            fg=RED,
            activebackground="#ffe5e5",
            relief="flat",
            padx=10,
            pady=5,
            font=("Arial", 8, "bold"),
            cursor="hand2",
        )
        self.export_button.pack(side="right")
        columns = ("contrato", "considerar", "desconsiderar", "status")
        self.table = ttk.Treeview(table_panel, columns=columns, show="headings", height=12)
        for key, label, width in (("contrato", "Contrato", 190), ("considerar", "Considerar", 135), ("desconsiderar", "Desconsiderar", 155), ("status", "Resultado", 200)):
            self.table.heading(key, text=label)
            self.table.column(key, width=width, anchor="w")
        self.table.tag_configure("odd", background="#fafafa")
        self.table.tag_configure("error", background="#fff0f0", foreground=RED)
        self.table.grid(row=1, column=0, sticky="nsew")
        scrollbar = ttk.Scrollbar(table_panel, orient="vertical", command=self.table.yview)
        scrollbar.grid(row=1, column=1, sticky="ns")
        self.table.configure(yscrollcommand=scrollbar.set)

    def _metric(self, parent, title, value, color):
        box = tk.Frame(parent, bg="#fafafa")
        tk.Label(box, text=title, fg=MUTED, bg="#fafafa", font=("Arial", 8, "bold")).pack(anchor="w")
        box.value_label = tk.Label(box, text=value, fg=color, bg="#fafafa", font=("Arial", 12, "bold"))
        box.value_label.pack(anchor="w", pady=(2, 0))
        return box

    def _choose_file(self):
        path = filedialog.askopenfilename(title="Selecione a planilha de contratos", filetypes=[("Planilhas Excel", "*.xlsx *.xls"), ("Todos os arquivos", "*.*")])
        if not path:
            return
        self.selected_file.set(path)
        try:
            primeira_linha = pd.read_excel(path, header=None, nrows=1)
            self.tem_cabecalho.set(detectar_cabecalho(primeira_linha))
            self._reload_columns()
        except Exception as error:
            self.selected_file.set("")
            messagebox.showerror("Nao foi possivel ler a planilha", str(error), parent=self)
            return
        self.selected_name.set(Path(path).name)
        self.selected_folder.set(str(Path(path).parent))
        self.remove_button.configure(state="normal")
        self.header_checkbox.configure(state="disabled" if Path(path).suffix.lower() in (".xlsx", ".xlsm") else "normal")
        if Path(path).suffix.lower() in (".xlsx", ".xlsm"):
            self.select_button.configure(state="disabled")
        if Path(path).suffix.lower() not in (".xlsx", ".xlsm"):
            self.status.set("Planilha carregada. Confira a coluna e a quantidade de contratos.")
        self._reset_results()

    def _excel_column_name(self, index):
        name = ""
        while index >= 0:
            index, remainder = divmod(index, 26)
            name = chr(65 + remainder) + name
            index -= 1
        return f"Coluna {name}"

    def _normalize_headers(self, values):
        names = []
        used = set()
        for index, value in enumerate(values):
            name = str(value).strip() if pd.notna(value) else ""
            if not name:
                name = self._excel_column_name(index)
            base = name
            suffix = 2
            while name.casefold() in used:
                name = f"{base} ({suffix})"
                suffix += 1
            used.add(name.casefold())
            names.append(name)
        return names

    def _reload_columns(self):
        path = self.selected_file.get()
        if not path:
            return
        try:
            if Path(path).suffix.lower() in (".xlsx", ".xlsm"):
                header = self.tem_cabecalho.get()
                self.column_box.configure(state="disabled")
                self.header_checkbox.configure(state="disabled")
                self.select_button.configure(state="disabled")
                self.remove_button.configure(state="disabled")
                self.status.set("Lendo a planilha em fluxo para contar os contratos...")
                threading.Thread(target=self._inspection_worker, args=(path, header), daemon=True).start()
                self.after(100, self._poll_inspection)
                return
            raw = pd.read_excel(path, header=None)
            if self.tem_cabecalho.get():
                names = self._normalize_headers(raw.iloc[0].tolist())
                self.preview_df = raw.iloc[1:].copy()
                self.preview_df.columns = names
                populated = [name for name in names if self.preview_df[name].notna().any()]
                preferred = next((name for name in populated if "contrato" in name.casefold()), populated[0] if populated else (names[0] if names else ""))
            else:
                self.preview_df = raw.copy()
                names = [self._excel_column_name(i) for i in range(len(raw.columns))]
                self.preview_df.columns = names
                populated = [name for name in names if self.preview_df[name].notna().any()]
                preferred = populated[0] if populated else (names[0] if names else "")
            self.columns = names
            self.column_counts = {}
            self.column_box.configure(values=names, state="readonly" if names else "disabled")
            self.contract_column.set(preferred)
            self._update_contract_count()
        except Exception as error:
            messagebox.showerror("Nao foi possivel ler a planilha", str(error), parent=self)

    def _inspection_worker(self, path, header):
        try:
            result = inspecionar_planilha_xlsx(path, header)
            self.inspection_queue.put((path, header, True, result))
        except Exception as error:
            self.inspection_queue.put((path, header, False, str(error)))

    def _poll_inspection(self):
        try:
            path, header, success, result = self.inspection_queue.get_nowait()
        except queue.Empty:
            self.after(100, self._poll_inspection)
            return
        if path != self.selected_file.get() or header != self.tem_cabecalho.get():
            return
        self.header_checkbox.configure(state="normal")
        self.select_button.configure(state="normal")
        self.remove_button.configure(state="normal" if self.selected_file.get() else "disabled")
        if not success:
            self.column_box.configure(state="disabled")
            self.status.set("Nao foi possivel ler a planilha.")
            messagebox.showerror("Nao foi possivel ler a planilha", result, parent=self)
            return
        self.preview_df = None
        self.columns = result["colunas"]
        self.column_counts = dict(zip(self.columns, result["contagens"]))
        populated = [name for name in self.columns if self.column_counts.get(name, 0)]
        preferred = next((name for name in populated if "contrato" in name.casefold()), populated[0] if populated else (self.columns[0] if self.columns else ""))
        self.column_box.configure(values=self.columns, state="readonly" if self.columns else "disabled")
        self.contract_column.set(preferred)
        self._update_contract_count()
        self.status.set("Planilha carregada. Confira a coluna e a quantidade de contratos.")

    def _update_contract_count(self, _event=None):
        column = self.contract_column.get()
        if not column or (column not in self.column_counts and (self.preview_df is None or column not in self.preview_df.columns)):
            self.contract_count.set("Contratos encontrados: --")
            return
        if column in self.column_counts:
            count = self.column_counts[column]
        else:
            values = self.preview_df[column]
            filled = values.notna() & values.astype(str).str.strip().ne("")
            count = int(filled.sum())
        label = "contrato encontrado" if count == 1 else "contratos encontrados"
        self.contract_count.set(f"{count} {label}")

    def _remove_file(self):
        if str(self.run_button.cget("state")) == "disabled":
            return
        if not messagebox.askyesno(
            "Confirmar remoção",
            f"Deseja remover '{self.selected_name.get()}' da tela?\n\nO arquivo original continuará salvo no computador.",
            parent=self,
        ):
            return
        self.selected_file.set("")
        self.selected_name.set("Nenhuma planilha selecionada")
        self.selected_folder.set("Excel .xlsx ou .xls")
        self.contract_column.set("")
        self.contract_count.set("Contratos encontrados: --")
        self.tem_cabecalho.set(True)
        self.preview_df = None
        self.column_counts = {}
        self.columns = []
        self.column_box.configure(values=(), state="disabled")
        self.header_checkbox.configure(state="disabled")
        self.remove_button.configure(state="disabled")
        self.status.set("Arquivo removido da selecao. O original continua no computador.")
        self._reset_results()

    def _reset_results(self):
        self.column_used.set("")
        self.log_rows = []
        self.details_db_path = None
        self.job_id = None
        self.log_page = 0
        self.report_metadata = None
        self.export_button.configure(state="disabled")
        for item in self.table.get_children():
            self.table.delete(item)
        self.total_label.value_label.configure(text="\u2014")
        self.success_label.value_label.configure(text="\u2014")
        self.fail_label.value_label.configure(text="\u2014")

    def _export_report(self):
        if not self.log_rows or not self.report_metadata:
            return

        original_file = self.report_metadata["arquivo"]
        default_name = f"Relatorio_DMA_{Path(original_file).stem}_{self.report_metadata['data_hora'].strftime('%Y%m%d_%H%M')}.xlsx"
        destination = filedialog.asksaveasfilename(
            title="Salvar relatório de processamento",
            initialfile=default_name,
            defaultextension=".xlsx",
            filetypes=[("Planilha Excel", "*.xlsx")],
        )
        if not destination:
            return

        summary = [
            ("Arquivo de origem", original_file),
            ("Data e hora", self.report_metadata["data_hora"].strftime("%d/%m/%Y %H:%M:%S")),
            ("Coluna de contratos", self.report_metadata["coluna"]),
            ("Cabeçalho na primeira linha", "Sim" if self.report_metadata["tem_cabecalho"] else "Não"),
            ("Ação DMA", self.report_metadata["acao"]),
            ("Ambiente", "API simulada" if self.report_metadata["modo"] == "simulacao" else "API real"),
            ("Total de contratos", self.report_metadata["total"]),
            ("Sucessos", self.report_metadata["sucesso"]),
            ("Falhas", self.report_metadata["falhas"]),
            ("Linhas vazias ignoradas", self.report_metadata["linhas_ignoradas"]),
        ]
        if self.job_id and self.details_db_path:
            try:
                self._export_report_stream(destination, summary)
                messagebox.showinfo("Relatório salvo", f"O relatório foi salvo em:\n{destination}", parent=self)
            except Exception as error:
                messagebox.showerror("Erro ao exportar relatório", str(error), parent=self)
            return

        details = [
            {
                "Contrato": safe_excel_text(item.get("contrato", "")),
                "Considerar DMA": str(item.get("considerar_dma", "")).lower(),
                "Desconsiderar DMA": str(item.get("desconsiderar_dma", "")).lower(),
                "Resultado": item.get("status", ""),
            }
            for item in self.log_rows
        ]

        try:
            from openpyxl.styles import Font, PatternFill

            with pd.ExcelWriter(destination, engine="openpyxl") as writer:
                pd.DataFrame(summary, columns=["Informação", "Valor"]).to_excel(writer, sheet_name="Resumo", index=False)
                pd.DataFrame(details).to_excel(writer, sheet_name="Contratos", index=False)

                for sheet in writer.book.worksheets:
                    sheet.freeze_panes = "A2"
                    for cell in sheet[1]:
                        cell.fill = PatternFill(fill_type="solid", fgColor="F40000")
                        cell.font = Font(color="FFFFFF", bold=True)
                    for column_cells in sheet.columns:
                        column_letter = column_cells[0].column_letter
                        max_width = max(len(str(cell.value or "")) for cell in column_cells)
                        sheet.column_dimensions[column_letter].width = min(max(max_width + 2, 14), 52)
                    if sheet.title == "Contratos":
                        sheet.auto_filter.ref = sheet.dimensions

            messagebox.showinfo("Relatório salvo", f"O relatório foi salvo em:\n{destination}", parent=self)
        except Exception as error:
            messagebox.showerror("Erro ao exportar relatório", str(error), parent=self)

    def _export_report_stream(self, destination, summary):
        from openpyxl import Workbook
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.styles import Font, PatternFill

        workbook = Workbook(write_only=True)
        header_fill = PatternFill(fill_type="solid", fgColor="F40000")
        header_font = Font(color="FFFFFF", bold=True)

        def append_header(sheet, labels):
            cells = []
            for label in labels:
                cell = WriteOnlyCell(sheet, value=label)
                cell.fill, cell.font = header_fill, header_font
                cells.append(cell)
            sheet.append(cells)

        summary_sheet = workbook.create_sheet("Resumo")
        append_header(summary_sheet, ["Informação", "Valor"])
        for row in summary:
            summary_sheet.append(list(row))
        summary_sheet.column_dimensions["A"].width = 32
        summary_sheet.column_dimensions["B"].width = 48

        details_sheet = workbook.create_sheet("Contratos")
        append_header(details_sheet, ["Contrato", "Considerar DMA", "Desconsiderar DMA", "Resultado"])
        for letter, width in (("A", 24), ("B", 20), ("C", 24), ("D", 72)):
            details_sheet.column_dimensions[letter].width = width

        database = connect_journal(self.details_db_path)
        try:
            for item in iter_items(database, self.job_id):
                details_sheet.append([
                    safe_excel_text(item["contract"]),
                    str(bool(item["consider_dma"])).lower(),
                    str(bool(item["disregard_dma"])).lower(),
                    item["status"],
                ])
        finally:
            database.close()
        workbook.save(destination)

    def _apply_log_filter(self, _event=None):
        for item in self.table.get_children():
            self.table.delete(item)

        filtro = self.log_filter.get()
        if self.job_id and self.details_db_path:
            database = connect_journal(self.details_db_path)
            try:
                count = count_items(database, self.job_id, filtro)
                total_pages = max(1, (count + self.page_size - 1) // self.page_size)
                self.log_page = min(self.log_page, total_pages - 1)
                visible_rows = page_items(database, self.job_id, filtro, self.page_size, self.log_page * self.page_size)
            finally:
                database.close()
        else:
            linhas = self.log_rows
            if filtro == "Somente falhas":
                linhas = [item for item in linhas if not str(item.get("status", "")).casefold().startswith("ok")]
            elif filtro == "Somente sucessos":
                linhas = [item for item in linhas if str(item.get("status", "")).casefold().startswith("ok")]
            count = len(linhas)
            total_pages = max(1, (count + self.page_size - 1) // self.page_size)
            self.log_page = min(self.log_page, total_pages - 1)
            start = self.log_page * self.page_size
            visible_rows = linhas[start : start + self.page_size]

        self.page_label.set(f"Página {self.log_page + 1} de {total_pages} ({count} linhas)")
        self.previous_page_button.configure(state="normal" if self.log_page else "disabled")
        self.next_page_button.configure(state="normal" if self.log_page + 1 < total_pages else "disabled")

        for index, item in enumerate(visible_rows):
            status = item["status"] if hasattr(item, "keys") else item.get("status", "")
            sucesso = str(status).casefold().startswith("ok")
            tags = () if sucesso else ("error",)
            if sucesso and index % 2:
                tags = ("odd",)
            self.table.insert(
                "",
                "end",
                values=(
                    item["contrato"] if hasattr(item, "keys") and "contrato" in item.keys() else item.get("contrato", ""),
                    str(item["consider_dma"] if hasattr(item, "keys") and "consider_dma" in item.keys() else item.get("considerar_dma", "")).lower(),
                    str(item["disregard_dma"] if hasattr(item, "keys") and "disregard_dma" in item.keys() else item.get("desconsiderar_dma", "")).lower(),
                    status,
                ),
                tags=tags,
            )

    def _previous_page(self):
        if self.log_page > 0:
            self.log_page -= 1
            self._apply_log_filter()

    def _change_log_filter(self, _event=None):
        self.log_page = 0
        self._apply_log_filter()

    def _next_page(self):
        self.log_page += 1
        self._apply_log_filter()

    def _process(self):
        path = self.selected_file.get()
        if not path:
            messagebox.showinfo("Selecione uma planilha", "Escolha um arquivo Excel antes de processar.", parent=self)
            return
        mode = "API simulada (nenhum dado será enviado à API real)" if MODO_TESTE_API else "API real"
        if not MODO_TESTE_API and (
            not CORAL_API_URL
            or not CORAL_API_URL.lower().startswith(("https://", "http://"))
            or "exemplo.com" in CORAL_API_URL.lower()
            or not get_coral_auth_token()
        ):
            messagebox.showerror(
                "Configuração incompleta",
                "Para usar a API real, configure CORAL_API_URL e CORAL_AUTH_TOKEN no arquivo .env.",
                parent=self,
            )
            return
        confirmation = (
            "Confirme os dados antes de iniciar:\n\n"
            f"Arquivo: {Path(path).name}\n"
            f"Coluna de contratos: {self.contract_column.get() or 'Não selecionada'}\n"
            f"Quantidade: {self.contract_count.get()}\n"
            f"Ação DMA: {self.action.get()}\n"
            f"Ambiente: {mode}\n\n"
            "Deseja continuar?"
        )
        if not messagebox.askyesno("Confirmar processamento", confirmation, parent=self):
            return
        try:
            if not Path(path).is_file():
                raise OSError("O arquivo selecionado não está mais disponível.")
            content = path
        except OSError as error:
            messagebox.showerror("Erro ao abrir arquivo", str(error), parent=self)
            return

        self.run_button.configure(state="disabled", text="Processando...")
        self.select_button.configure(state="disabled")
        self.remove_button.configure(state="disabled")
        self.column_box.configure(state="disabled")
        self.header_checkbox.configure(state="disabled")
        self.disregard_radio.configure(state="disabled")
        self.consider_radio.configure(state="disabled")
        self.status.set("Processando na API simulada..." if MODO_TESTE_API else "Enviando contratos para a API real...")
        self._reset_results()
        threading.Thread(target=self._worker, args=(content, self.action.get(), self.contract_column.get() or None, self.tem_cabecalho.get(), False), daemon=True).start()
        self.after(100, self._poll_result)

    def _worker(self, content, action, column, tem_cabecalho, force_uncertain=False, force_new=False):
        try:
            self.result_queue.put((True, processar_e_enviar_planilha(
                content, action, column, tem_cabecalho,
                force_uncertain=force_uncertain, stream_results=True, force_new=force_new,
            )))
        except Exception as error:
            self.result_queue.put((False, str(error)))

    def _start_followup(self, force_uncertain=False, force_new=False):
        path = self.selected_file.get()
        self._reset_results()
        self.run_button.configure(state="disabled", text="Retomando...")
        self.select_button.configure(state="disabled")
        self.remove_button.configure(state="disabled")
        self.column_box.configure(state="disabled")
        self.header_checkbox.configure(state="disabled")
        self.disregard_radio.configure(state="disabled")
        self.consider_radio.configure(state="disabled")
        self.status.set("Reenviando somente após confirmação..." if force_uncertain else "Iniciando um novo processamento...")
        threading.Thread(
            target=self._worker,
            args=(path, self.action.get(), self.contract_column.get() or None, self.tem_cabecalho.get(), force_uncertain, force_new),
            daemon=True,
        ).start()
        self.after(100, self._poll_result)

    def _poll_result(self):
        try:
            success, result = self.result_queue.get_nowait()
        except queue.Empty:
            self.after(100, self._poll_result)
            return

        self.run_button.configure(state="normal", text="Processar contratos  \u2192")
        self.select_button.configure(state="normal")
        self.remove_button.configure(state="normal" if self.selected_file.get() else "disabled")
        self.column_box.configure(state="readonly" if self.columns else "disabled")
        self.header_checkbox.configure(state="normal" if self.selected_file.get() else "disabled")
        self.disregard_radio.configure(state="normal")
        self.consider_radio.configure(state="normal")
        if not success:
            self.status.set("Nao foi possivel processar a planilha.")
            messagebox.showerror("Erro no processamento", result, parent=self)
            return

        self.details_db_path = result.get("detalhes_db")
        self.job_id = result.get("job_id")
        self.total_label.value_label.configure(text=str(result["total"]))
        self.success_label.value_label.configure(text=str(result["sucesso"]))
        self.fail_label.value_label.configure(text=str(result["falhas"]))
        simulated = result.get("modo") == "simulacao"
        self.status.set("API simulada. Nenhum dado enviado para a API real." if simulated else "Processamento concluido pela API real.")
        ignored = result.get("linhas_ignoradas", 0)
        tipo_coluna = result.get("coluna_contrato", "")
        prefixo_coluna = "" if result.get("tem_cabecalho") else " (sem cabecalho)"
        self.column_used.set(f"Coluna usada: {tipo_coluna}{prefixo_coluna}. {ignored} linha(s) vazia(s) ignorada(s).")
        self.log_rows = result.get("detalhes") or []
        self.report_metadata = {
            "arquivo": self.selected_name.get(),
            "data_hora": datetime.now().astimezone(),
            "coluna": tipo_coluna,
            "tem_cabecalho": result.get("tem_cabecalho", False),
            "acao": self.action.get(),
            "modo": result.get("modo", "simulacao"),
            "total": result.get("total", 0),
            "sucesso": result.get("sucesso", 0),
            "falhas": result.get("falhas", 0),
            "linhas_ignoradas": ignored,
        }
        self.export_button.configure(state="normal" if self.log_rows or self.job_id else "disabled")
        if result.get("incertos", 0) and not result.get("forcou_incertos"):
            resend = messagebox.askyesno(
                "Há resultados incertos",
                f"{result['incertos']} contrato(s) podem ter sido alterados pela Coral, mas a confirmação não chegou. Confira esses contratos na Coral antes de decidir.\n\nDeseja reenviar mesmo assim?",
                parent=self,
            )
            if resend:
                self._start_followup(force_uncertain=True)
                return
        if result.get("ja_processados", 0) and not result.get("enviados", 0) and not result.get("forcou_novo"):
            rerun = messagebox.askyesno(
                "Contratos já processados",
                f"Todos os {result['ja_processados']} contratos já aparecem como confirmados neste processamento. Deseja iniciar uma nova execução e enviá-los novamente?",
                parent=self,
            )
            if rerun:
                self._start_followup(force_new=True)
                return
        self._apply_log_filter()


if __name__ == "__main__":
    DMADesktopApp().mainloop()
