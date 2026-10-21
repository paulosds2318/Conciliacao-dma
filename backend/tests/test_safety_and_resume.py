import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

from openpyxl import Workbook
from requests import Timeout

from app.pipeline import processar_e_enviar_planilha
from app.spreadsheet import inspecionar_planilha_xlsx, iterar_contratos_xlsx, normalizar_cabecalhos
from app.journal import connect, count_items, page_items


class SpreadsheetTests(unittest.TestCase):
    def test_preserves_text_ids_and_formatted_numeric_leading_zeroes(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "contracts.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Contrato", "Outro"])
            sheet.append(["000012345678", "x"])
            sheet.append([1234, "y"])
            sheet["A3"].number_format = "000000"
            workbook.save(path)

            info = inspecionar_planilha_xlsx(path, True)
            self.assertEqual(info["contagens"], [2, 2])
            self.assertEqual(list(iterar_contratos_xlsx(path, "Contrato", True)), ["000012345678", "001234"])

    def test_normalizes_empty_and_duplicate_headers(self):
        self.assertEqual(normalizar_cabecalhos([None, "Contrato", "Contrato", ""]), [
            "Coluna A", "Contrato", "Contrato (2)", "Coluna D"
        ])


class ResumeTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        os.environ["CONCILIACAO_DMA_DATA_DIR"] = self.temporary.name
        self.addCleanup(os.environ.pop, "CONCILIACAO_DMA_DATA_DIR", None)
        self.path = Path(self.temporary.name) / "input.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Contrato"])
        sheet.append(["00000001"])
        sheet.append(["00000002"])
        workbook.save(self.path)

    @patch("app.pipeline.get_coral_auth_token", return_value="test-token")
    @patch("app.pipeline.MODO_TESTE_API", False)
    @patch("app.pipeline.CORAL_API_URL", "https://api.test.invalid/dma")
    @patch("app.pipeline._logger")
    @patch("app.pipeline.requests.post")
    def test_successful_contracts_are_not_sent_twice(self, post, _logger, _token):
        post.return_value = Mock(status_code=200)
        first = processar_e_enviar_planilha(self.path, "Considerar DMA", "Contrato", True, stream_results=True)
        second = processar_e_enviar_planilha(self.path, "Considerar DMA", "Contrato", True, stream_results=True)

        self.assertEqual(first["sucesso"], 2)
        self.assertIsNone(first["detalhes"])
        self.assertEqual(second["sucesso"], 2)
        self.assertEqual(post.call_count, 2)
        database = connect(first["detalhes_db"])
        try:
            self.assertEqual(count_items(database, first["job_id"]), 2)
            self.assertEqual(len(page_items(database, first["job_id"], limit=1)), 1)
        finally:
            database.close()
        from types import SimpleNamespace
        report = Path(self.temporary.name) / "report.xlsx"
        DMADesktopApp._export_report_stream(
            SimpleNamespace(details_db_path=first["detalhes_db"], job_id=first["job_id"]),
            report,
            [("Total", 2)],
        )
        from openpyxl import load_workbook
        workbook = load_workbook(report, read_only=True, data_only=True)
        rows = list(workbook["Contratos"].iter_rows(values_only=True))
        workbook.close()
        self.assertEqual(len(rows), 3)

    @patch("app.pipeline.get_coral_auth_token", return_value="test-token")
    @patch("app.pipeline.MODO_TESTE_API", False)
    @patch("app.pipeline.CORAL_API_URL", "https://api.test.invalid/dma")
    @patch("app.pipeline._logger")
    @patch("app.pipeline.requests.post")
    def test_uncertain_send_requires_explicit_confirmation_to_retry(self, post, _logger, _token):
        post.side_effect = [Timeout(), Mock(status_code=200), Mock(status_code=200)]
        first = processar_e_enviar_planilha(self.path, "Desconsiderar DMA", "Contrato", True, stream_results=True)
        resumed = processar_e_enviar_planilha(self.path, "Desconsiderar DMA", "Contrato", True, stream_results=True)

        self.assertEqual(first["incertos"], 1)
        self.assertEqual(resumed["incertos"], 1)
        self.assertEqual(post.call_count, 2)

        confirmed = processar_e_enviar_planilha(
            self.path, "Desconsiderar DMA", "Contrato", True,
            force_uncertain=True, stream_results=True,
        )
        self.assertEqual(confirmed["sucesso"], 2)
        self.assertEqual(post.call_count, 3)


if __name__ == "__main__":
    unittest.main()
